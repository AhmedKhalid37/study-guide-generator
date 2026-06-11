from __future__ import annotations

import csv
import re
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pipeline.extraction_metadata import classify_pdf_page_record
from pipeline.ocr_provider import OcrRequest, get_default_ocr_provider
from pipeline.ocr_routing import decide_ocr_route


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".docx",
    ".pptx",
    ".pdf",
}


class ExtractionError(RuntimeError):
    pass


class PdfEncryptedError(ExtractionError):
    """A PDF needs a password before its pages can be read.

    Raised only by :func:`preflight_pdf`; the existing extraction path is
    unchanged. Lets the API distinguish an *encrypted* PDF (clear "unlock it"
    message) from a generically *corrupt* one.
    """


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    mode: str
    warnings: list[str]
    metadata: dict[str, Any] | None = None


@dataclass(frozen=True)
class PdfPreflightResult:
    """Cheap, read-only inspection of a PDF (no OCR, no full extraction).

    Produced by :func:`preflight_pdf` from a bounded page sample so the API can
    warn about large/scanned PDFs *before* a job is created. Counts derived from
    ``text_ratio`` are estimates extrapolated from the sample, not exact.
    """

    page_count: int
    sampled_pages: int
    meaningful_pages: int
    text_ratio: float
    image_ratio: float
    scanned_flag: str  # "text" | "mixed" | "image_heavy" | "unknown"
    ocr_available: bool


def extract_file(path: Path, pages: Iterable[int] | None = None) -> ExtractionResult:
    """Extract text from a supported attachment.

    ``pages`` is an optional set of **1-based, original** page numbers to restrict
    extraction to; it applies **only to PDFs** (it mirrors the per-PDF
    ``page_selections`` request field). For every other file type it is ignored —
    page selections are meaningless for non-paginated sources. ``pages=None`` (the
    default) preserves the exact previous behaviour for all types.
    """
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(f"Unsupported attachment type: {suffix or 'unknown'}")

    if suffix in {".txt", ".md", ".markdown"}:
        return ExtractionResult(path.read_text(encoding="utf-8", errors="replace"), "text", [])
    if suffix in {".csv", ".tsv"}:
        return ExtractionResult(_extract_table(path, delimiter="\t" if suffix == ".tsv" else ","), "table", [])
    if suffix == ".docx":
        return ExtractionResult(_extract_docx(path), "docx", [])
    if suffix == ".pptx":
        return ExtractionResult(_extract_pptx(path), "pptx", [])
    if suffix == ".pdf":
        return _extract_pdf(path, pages=pages)

    raise ExtractionError(f"Unsupported attachment type: {suffix or 'unknown'}")


def _extract_table(path: Path, *, delimiter: str) -> str:
    rows: list[str] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for row in reader:
            rows.append(" | ".join(cell.strip() for cell in row))
    return "\n".join(rows)


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ExtractionError("python-docx is not installed.") from exc

    try:
        document = Document(path)
    except Exception as exc:
        raise ExtractionError(
            f"the file may be corrupt or unreadable ({type(exc).__name__})"
        ) from exc
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ExtractionError("python-pptx is not installed.") from exc

    try:
        presentation = Presentation(path)
    except Exception as exc:
        raise ExtractionError(
            f"the file may be corrupt or unreadable ({type(exc).__name__})"
        ) from exc
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        slide_parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())
        if slide_parts:
            parts.append(f"Slide {index}\n" + "\n".join(slide_parts))
    return "\n\n".join(parts)


def _extract_pdf(path: Path, pages: Iterable[int] | None = None) -> ExtractionResult:
    """Extract a PDF page-by-page, optionally restricted to selected pages.

    ``pages`` is an optional iterable of **1-based, original** page numbers. When
    ``None`` (the default) the whole document is processed exactly as before — this
    path is byte-for-byte unchanged. When provided, only those pages are read, but
    each kept page keeps its **original** ``## Page N`` anchor (so selecting pages
    20-21 yields ``## Page 20`` / ``## Page 21``, never renumbered to 1-2). The
    per-page text-vs-OCR fallback and the ``pdf_text``/``pdf_ocr``/``pdf_mixed``
    mode reporting are computed over the selected subset, unchanged otherwise.

    Selected pages are validated against the real page count: out-of-range pages
    are ignored with a warning, never raising. If a selection is given but no page
    is in range, an empty result with a clear warning is returned (the caller then
    reports "no text could be extracted", matching existing behaviour).
    """
    try:
        import fitz
    except ImportError as exc:
        raise ExtractionError("PyMuPDF is not installed.") from exc

    warnings: list[str] = []
    try:
        document_ctx = fitz.open(path)
    except Exception as exc:
        raise ExtractionError(
            f"the file may be corrupt or an unreadable scan ({type(exc).__name__})"
        ) from exc
    with document_ctx as document:
        page_count = int(document.page_count)
        # Resolve the optional page selection against the real page count. Pages
        # are 1-based and refer to ORIGINAL page numbers; out-of-range entries are
        # dropped (not clamped) so a stray "page 999" can't pull in the last page,
        # and so anchors stay truthful. None => no filter => process every page.
        selected: set[int] | None = None
        if pages is not None:
            requested = {int(p) for p in pages}
            selected = {p for p in requested if 1 <= p <= page_count}
            out_of_range = sorted(requested - selected)
            if out_of_range:
                warnings.append(
                    f"Ignored selected page(s) outside this {page_count}-page PDF: "
                    f"{', '.join(str(p) for p in out_of_range)}."
                )
            if not selected:
                warnings.append(
                    f"None of the selected pages fall within this {page_count}-page PDF; "
                    "no pages were extracted."
                )
                return ExtractionResult(
                    "",
                    "pdf_text",
                    warnings,
                    _pdf_metadata(page_count=page_count, pages=[], warnings=warnings),
                )
        # Prefix each page with a "## Page N" anchor (mirroring the pptx "Slide N"
        # marker) so positional references survive extraction — study-guide prompts
        # cite these. The index is the physical page number; blank pages are dropped
        # without shifting the numbering of the pages that follow.
        #
        # Page-level (not whole-document) text/OCR fallback: a mostly-scanned PDF
        # whose title slide happens to carry a few words of embedded text must not
        # be treated as "fully text-extracted". Each page decides independently —
        # pages with meaningful embedded text use it; sparse/scanned pages get OCR'd
        # on their own — so a 118-page scanned deck yields all its pages, not just
        # page 1.
        ocr_provider = get_default_ocr_provider()
        ocr_ready, ocr_unavailable_reason = ocr_provider.is_available()

        blocks: list[str] = []
        page_metadata: list[dict[str, Any]] = []
        used_text = False
        used_ocr = False
        wanted_ocr = False
        for index, page in enumerate(document, start=1):
            # When a selection is active, skip unselected pages before any text or
            # OCR work — so OCR only ever runs on selected pages — while `index`
            # stays the original page number for the anchor below.
            if selected is not None and index not in selected:
                continue
            # Advisory-only visual/object signals (image/drawing counts, page
            # dimensions). Collected once per processed page; never changes the
            # extracted text or which path a page takes. Degrades to unknown-safe
            # values on any failure.
            visual = _pdf_visual_signals(page)
            body = page.get_text("text").strip()
            if _is_meaningful_page_text(body):
                blocks.append(f"## Page {index}\n{body}")
                page_metadata.append(
                    _pdf_page_metadata(
                        index, "embedded_text", body,
                        has_page_anchor=True, visual=visual, ocr_ready=ocr_ready,
                    )
                )
                used_text = True
                continue

            # Sparse/empty page → OCR it individually. A page uses embedded text
            # XOR OCR (never both), so the same content is not duplicated.
            wanted_ocr = True
            page_ocr = (
                ocr_provider.ocr_page(OcrRequest(page=page, page_number=index)).text
                if ocr_ready
                else ""
            )
            if page_ocr:
                blocks.append(f"## Page {index}\n{page_ocr}")
                page_metadata.append(
                    _pdf_page_metadata(
                        index, "ocr", page_ocr,
                        has_page_anchor=True, visual=visual, ocr_ready=ocr_ready,
                    )
                )
                used_ocr = True
            elif body:
                # OCR produced nothing (or is unavailable) but the page had a little
                # embedded text — keep it rather than dropping the page entirely.
                blocks.append(f"## Page {index}\n{body}")
                page_warnings = ["ocr_unavailable"] if not ocr_ready else ["ocr_empty_fallback_embedded_text"]
                page_metadata.append(
                    _pdf_page_metadata(
                        index,
                        "embedded_text",
                        body,
                        has_page_anchor=True,
                        warnings=page_warnings,
                        visual=visual,
                        ocr_ready=ocr_ready,
                    )
                )
                used_text = True
            else:
                # Genuinely blank/unreadable page -> dropped from text output
                # (numbering preserved by index), but still represented in the
                # audit metadata.
                page_warnings = ["ocr_unavailable"] if not ocr_ready else ["no_text_extracted"]
                page_metadata.append(
                    _pdf_page_metadata(
                        index,
                        "none",
                        "",
                        has_page_anchor=False,
                        warnings=page_warnings,
                        visual=visual,
                        ocr_ready=ocr_ready,
                    )
                )

        if wanted_ocr and not ocr_ready and ocr_unavailable_reason:
            warnings.append(ocr_unavailable_reason)

        text = "\n\n".join(blocks)
        if used_ocr and used_text:
            mode = "pdf_mixed"
        elif used_ocr:
            mode = "pdf_ocr"
        elif used_text:
            mode = "pdf_text"
        else:
            # No usable text at all: OCR was intended but produced nothing
            # (unavailable, or genuinely empty scans). Report the OCR mode as the
            # all-or-nothing path did, so downstream "no text extracted" handling
            # is unchanged.
            mode = "pdf_ocr" if wanted_ocr else "pdf_text"
        return ExtractionResult(
            text,
            mode,
            warnings,
            _pdf_metadata(page_count=page_count, pages=page_metadata, warnings=warnings),
        )


def _pdf_metadata(
    *,
    page_count: int,
    pages: list[dict[str, Any]],
    warnings: list[str],
) -> dict[str, Any]:
    return {
        "kind": "pdf_extraction",
        "page_count": int(page_count),
        "pages": pages,
        "warnings": list(warnings),
    }


def _pdf_page_metadata(
    page: int,
    method: str,
    text: str,
    *,
    has_page_anchor: bool,
    warnings: list[str] | None = None,
    visual: dict[str, Any] | None = None,
    ocr_ready: bool = False,
) -> dict[str, Any]:
    record: dict[str, Any] = {
        "page": int(page),
        "method": method,
        "text_chars": len(text),
        "word_count": len(re.findall(r"\w+", text)),
        "has_page_anchor": bool(has_page_anchor),
        "warnings": list(warnings or []),
    }
    # Additive, advisory-only visual/object signals (Slice 30). Always merged so
    # the per-page schema is uniform; values are unknown-safe (None / null) when a
    # signal could not be collected. Never affects the legacy fields above.
    if visual:
        record.update(visual)
    # Advisory OCR routing decision (Slice 34). Recorded once the page's method,
    # text and visual signals are known; `ocr_ready` is the document-level local
    # OCR availability the extractor already probed. This is a *recorder, not a
    # router*: it never decides whether OCR actually runs (that decision is the
    # unchanged per-page logic in `_extract_pdf`), so extracted text is identical.
    record.update(_pdf_route_decision(record, ocr_ready=ocr_ready))
    return record


def _pdf_route_decision(record: dict[str, Any], *, ocr_ready: bool) -> dict[str, Any]:
    """Advisory, local-only OCR routing decision for one page (Slice 34).

    Maps the page's already-collected signals to the pure routing policy
    (:func:`pipeline.ocr_routing.decide_ocr_route`) via its advisory
    ``classification`` (shared with the persisted artifact through
    :func:`classify_pdf_page_record`, so both agree). Local-only: cloud OCR stays
    off, so the provider is only ever ``tesseract_local`` or ``None``.

    Degrade-not-fail: the policy itself never raises, but the whole adapter is
    wrapped so that any unexpected error records a fixed safe route instead of
    breaking extraction. Returns only closed-vocabulary route tokens — never a
    path, secret, image blob, or free-text error.
    """
    try:
        classification = classify_pdf_page_record(record).get("classification")
        decision = decide_ocr_route(
            {"classification": classification},
            # Local-first, cloud-off (the policy default); only the live local-OCR
            # availability is overridden so a `likely_scanned` page routes to local
            # OCR exactly when the extractor itself could OCR it.
            {"local_ocr_available": bool(ocr_ready), "allow_cloud_ocr": False},
        )
        return {
            "ocr_route_action": decision.get("action"),
            "ocr_route_provider": decision.get("provider"),
            "ocr_route_reason": decision.get("reason"),
            "ocr_route_confidence": decision.get("confidence"),
            "ocr_route_warnings": list(decision.get("warnings") or []),
        }
    except Exception:
        # Routing must never fail a generation; fall back to a fixed safe route.
        return {
            "ocr_route_action": "unknown",
            "ocr_route_provider": None,
            "ocr_route_reason": "routing_unavailable",
            "ocr_route_confidence": "low",
            "ocr_route_warnings": ["routing_input_unrecognized"],
        }


def _pdf_visual_signals(page) -> dict[str, Any]:
    """Cheap, additive per-page visual/object signals for future OCR routing.

    Advisory-only and **never raises**. Uses PyMuPDF read-only inspections
    (``page.get_images()`` / ``page.get_drawings()`` / ``page.rect``) that do not
    rasterise the page and do not change the extracted text or which extraction
    path a page takes. Returns numeric/boolean signals only — no image bytes,
    object data, file paths, or text.

    Each signal degrades independently: if a PyMuPDF method is missing or a single
    call fails, that field falls back to an unknown-safe value (``None`` count,
    ``None`` bool) and a short category is recorded in ``visual_warnings`` (no
    exception text, no paths). Extraction itself is unaffected.
    """
    image_count: int | None
    drawing_count: int | None
    has_images: bool | None
    has_drawings: bool | None
    page_width: float | None
    page_height: float | None
    visual_warnings: list[str] = []

    try:
        images = page.get_images(full=False)
        image_count = len(images) if images is not None else 0
        has_images = image_count > 0
    except Exception:
        image_count = None
        has_images = None
        visual_warnings.append("visual_image_signal_unavailable")

    try:
        drawings = page.get_drawings()
        drawing_count = len(drawings) if drawings is not None else 0
        has_drawings = drawing_count > 0
    except Exception:
        drawing_count = None
        has_drawings = None
        visual_warnings.append("visual_drawing_signal_unavailable")

    try:
        rect = page.rect
        page_width = round(float(rect.width), 2)
        page_height = round(float(rect.height), 2)
    except Exception:
        page_width = None
        page_height = None
        visual_warnings.append("visual_dimension_signal_unavailable")

    signals: dict[str, Any] = {
        "image_object_count": image_count,
        "drawing_object_count": drawing_count,
        "has_images": has_images,
        "has_drawings": has_drawings,
        "page_width": page_width,
        "page_height": page_height,
    }
    if visual_warnings:
        signals["visual_warnings"] = visual_warnings
    return signals


def _is_meaningful_page_text(text: str) -> bool:
    """Whether a page's embedded text is rich enough to skip OCR for that page.

    Conservative gate for the mixed scanned/text PDF path: a page counts as a real
    text page when it has either a reasonable amount of characters (>= 40) or
    several word-like tokens (>= 5). Sparse pages — blank scans, or image-only
    slides carrying just a stray label — fall through to per-page OCR instead of
    being accepted as "text extracted".
    """
    stripped = text.strip()
    if len(stripped) >= 40:
        return True
    return len(re.findall(r"\w+", stripped)) >= 5


def _ocr_available() -> tuple[bool, str | None]:
    """Backward-compatible OCR availability probe.

    Slice 32 moved the OCR engine behind the :mod:`pipeline.ocr_provider` boundary;
    this thin shim delegates to the default provider so the ``(ready, reason)``
    shape and the exact user-facing reason strings are unchanged. Retained for
    ``preflight_pdf`` and existing tests that import it directly.
    """
    return get_default_ocr_provider().is_available()


def _preflight_sample_indices(page_count: int, sample_pages: int) -> list[int]:
    """Bounded, evenly spaced 0-based page indices to probe for the estimate.

    Returns every page when the document is small enough; otherwise spreads
    ``sample_pages`` probes across the whole document (always including the first
    page) so the text-vs-scanned ratio is representative without reading every
    page. Cheap by construction — at most ``sample_pages`` ``get_text`` calls.
    """
    if page_count <= 0 or sample_pages <= 0:
        return []
    if page_count <= sample_pages:
        return list(range(page_count))
    step = page_count / sample_pages
    return sorted({min(page_count - 1, int(i * step)) for i in range(sample_pages)})


def _scanned_flag(image_ratio: float, *, image_heavy_ratio: float, text_ratio: float) -> str:
    if image_ratio <= text_ratio:
        return "text"
    if image_ratio >= image_heavy_ratio:
        return "image_heavy"
    return "mixed"


def preflight_pdf(
    path: Path,
    *,
    sample_pages: int = 20,
    image_heavy_ratio: float = 0.85,
    text_ratio: float = 0.15,
) -> PdfPreflightResult:
    """Inspect a PDF cheaply: page count + sampled text-vs-scanned estimate.

    Read-only and OCR-free. Opens the PDF with PyMuPDF (``fitz``), reads the page
    count, then probes a bounded sample of pages with ``page.get_text("text")``
    gated by the *same* :func:`_is_meaningful_page_text` heuristic the real
    extractor uses — so the estimate tracks actual extraction behavior. Does
    **not** touch :func:`_extract_pdf`, never rasterizes/OCRs, and persists
    nothing.

    Raises :class:`PdfEncryptedError` for password-protected PDFs and
    :class:`ExtractionError` for a corrupt/unopenable file (mirroring
    ``_extract_pdf``'s corrupt message). ``import fitz`` propagating ``ImportError``
    lets the caller degrade gracefully when PyMuPDF is unavailable.
    """
    import fitz  # may raise ImportError -> caller degrades to a safe verdict

    try:
        document_ctx = fitz.open(path)
    except Exception as exc:
        raise ExtractionError(
            f"the file may be corrupt or not a real PDF ({type(exc).__name__})"
        ) from exc

    with document_ctx as document:
        if getattr(document, "needs_pass", False):
            raise PdfEncryptedError("the PDF is encrypted or password-protected")

        page_count = int(document.page_count)
        indices = _preflight_sample_indices(page_count, sample_pages)
        meaningful = 0
        for index in indices:
            body = document[index].get_text("text").strip()
            if _is_meaningful_page_text(body):
                meaningful += 1
        sampled = len(indices)
        ocr_ready, _ = _ocr_available()

    if sampled == 0:
        return PdfPreflightResult(
            page_count=page_count,
            sampled_pages=0,
            meaningful_pages=0,
            text_ratio=0.0,
            image_ratio=0.0,
            scanned_flag="unknown",
            ocr_available=ocr_ready,
        )

    text_ratio_val = meaningful / sampled
    image_ratio = 1.0 - text_ratio_val
    flag = _scanned_flag(image_ratio, image_heavy_ratio=image_heavy_ratio, text_ratio=text_ratio)
    return PdfPreflightResult(
        page_count=page_count,
        sampled_pages=sampled,
        meaningful_pages=meaningful,
        text_ratio=round(text_ratio_val, 4),
        image_ratio=round(image_ratio, 4),
        scanned_flag=flag,
        ocr_available=ocr_ready,
    )
