"""Convert a cleaned Markdown file into an editable .docx (Word) document.

Source of truth is the job's ``clean.md`` — the same sanitized Markdown the PDF
and HTML are rendered from. This module never re-runs the LLM and never touches
the PDF/HTML render path; it is a standalone, dependency-light converter built on
``markdown-it-py`` (already used by the HTML renderer) + ``python-docx``.

Document fidelity:
  - a centered cover page with the guide title + generation date, then a page break
  - a running header on every page (guide title, small, top-right)
  - a footer with centered "Page X of Y" page numbers
  - headings (h1-h6 -> Word Heading styles)
  - paragraphs with **bold**, *italic*, ~~strikethrough~~, `inline code`, links
  - bullet and numbered lists (incl. one level of nesting)
  - fenced/indented code blocks (monospace, light grey background, indented)
  - tables with visible borders and a bold, shaded header row
  - inline images embedded from disk (scaled to fit the page)
  - horizontal rules

Known limitations (documented, not silent):
  - Math is preserved as literal LaTeX text (``$...$`` / ``$$...$$``), not rendered
    as equations — DOCX has no KaTeX equivalent in this pipeline.
  - An image that is missing on disk (or an unsupported/remote source) leaves a
    small ``[image missing: name]`` marker rather than failing the export.
  - Deeply nested lists collapse to the deepest supported Word list style (3).
  - Raw HTML in Markdown is rendered as plain text, not interpreted.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from markdown_it import MarkdownIt

# Mirror the HTML renderer's Markdown dialect so DOCX and HTML agree on structure
# (tables + strikethrough enabled). html:False — we never interpret raw HTML here.
_PARSER = (
    MarkdownIt("commonmark", {"html": False, "linkify": True, "typographer": False})
    .enable("table")
    .enable("strikethrough")
)

_CODE_FONT = "Courier New"
_CODE_SHADE = "F2F2F2"  # light grey background for code blocks
_HEADER_SHADE = "E7E7E7"  # light grey shading for table header row
_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".tif"}


def render_docx(markdown_path, output_path, *, title=None, generated_date=None):
    """Render ``markdown_path`` (a clean.md) into a .docx at ``output_path``.

    ``title``/``generated_date`` populate the cover page and running header. When
    omitted they fall back to the first H1 in the document and today's date.
    """
    from docx import Document

    source = Path(markdown_path)
    target = Path(output_path)
    text = source.read_text(encoding="utf-8", errors="replace")
    tokens = _PARSER.parse(text)

    document = Document()
    doc_title = (title or "").strip() or _first_heading(tokens) or "Study Guide"
    date_label = _format_date(generated_date)

    _add_cover_page(document, doc_title, date_label)
    _configure_running_header_footer(document, doc_title)
    _emit_blocks(document, tokens, base_dir=source.parent)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))
    return target


def _first_heading(tokens) -> str:
    for i, token in enumerate(tokens):
        if token.type == "heading_open" and i + 1 < len(tokens):
            inline = tokens[i + 1]
            if inline.type == "inline":
                return "".join(
                    child.content
                    for child in (inline.children or [])
                    if child.type in ("text", "code_inline")
                ).strip()
    return ""


def _format_date(generated_date) -> str:
    if generated_date:
        raw = str(generated_date)
        for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M", "%Y-%m-%d"):
            try:
                return datetime.strptime(raw[: len(fmt) + 2], fmt).strftime("%B %d, %Y")
            except ValueError:
                continue
        # Not an ISO timestamp we recognise — show whatever the caller passed.
        return raw
    return datetime.now().strftime("%B %d, %Y")


def _add_cover_page(document, title: str, date_label: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    # Push the title toward vertical center with a little leading space.
    for _ in range(6):
        document.add_paragraph()

    title_p = document.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(32)

    date_p = document.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_p.add_run(date_label)
    date_run.font.size = Pt(14)
    date_run.italic = True

    _add_page_break(document)


def _add_page_break(document) -> None:
    from docx.enum.text import WD_BREAK

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _configure_running_header_footer(document, title: str) -> None:
    """Title in a small top-right header; centered "Page X of Y" footer.

    Uses a distinct first-page header/footer (left blank) so the cover page stays
    clean while every content page carries the running header and page numbers.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    section = document.sections[0]
    section.different_first_page_header_footer = True

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_p.add_run(title)
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = _grey()

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("Page ")
    run.font.size = Pt(9)
    _add_field(footer_p, "PAGE", size_pt=9)
    mid = footer_p.add_run(" of ")
    mid.font.size = Pt(9)
    _add_field(footer_p, "NUMPAGES", size_pt=9)


def _grey():
    from docx.shared import RGBColor

    return RGBColor(0x6B, 0x6B, 0x6B)


def _add_field(paragraph, field_code: str, *, size_pt: int | None = None) -> None:
    """Append a Word field (e.g. PAGE / NUMPAGES) to ``paragraph``."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run = paragraph.add_run()
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _emit_blocks(document, tokens, *, base_dir: Path) -> None:
    list_stack: list[str] = []  # "bullet" | "ordered", deepest last
    i = 0
    n = len(tokens)
    while i < n:
        token = tokens[i]
        kind = token.type

        if kind == "heading_open":
            level = min(int(token.tag[1:] or 1), 9)
            paragraph = document.add_heading(level=level)
            _emit_inline(document, paragraph, tokens[i + 1], base_dir=base_dir)
            i += 3  # heading_open, inline, heading_close
            continue

        if kind == "paragraph_open":
            inline = tokens[i + 1]
            if list_stack:
                paragraph = document.add_paragraph(style=_list_style(list_stack[-1], len(list_stack)))
            else:
                paragraph = document.add_paragraph()
            _emit_inline(document, paragraph, inline, base_dir=base_dir)
            i += 3  # paragraph_open, inline, paragraph_close
            continue

        if kind == "bullet_list_open":
            list_stack.append("bullet")
            i += 1
            continue
        if kind == "ordered_list_open":
            list_stack.append("ordered")
            i += 1
            continue
        if kind in ("bullet_list_close", "ordered_list_close"):
            if list_stack:
                list_stack.pop()
            i += 1
            continue
        if kind in ("list_item_open", "list_item_close"):
            i += 1
            continue

        if kind in ("fence", "code_block"):
            _emit_code_block(document, token.content)
            i += 1
            continue

        if kind == "hr":
            document.add_paragraph("—" * 24)
            i += 1
            continue

        if kind == "table_open":
            i = _emit_table(document, tokens, i, base_dir=base_dir)
            continue

        # blockquote_open/close, html_block, and any unhandled wrappers: skip the
        # marker; their inner inline/paragraph tokens are emitted by the rules above.
        i += 1


def _list_style(kind: str, depth: int) -> str:
    """Map a list kind + nesting depth to a Word built-in list style name."""
    base = "List Bullet" if kind == "bullet" else "List Number"
    if depth <= 1:
        return base
    return f"{base} {min(depth, 3)}"


def _emit_inline(document, paragraph, inline_token, *, base_dir: Path) -> None:
    if inline_token is None or inline_token.type != "inline":
        return
    bold = italic = strike = False
    for child in inline_token.children or []:
        ctype = child.type
        if ctype == "text":
            run = paragraph.add_run(child.content)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if strike:
                run.font.strike = True
        elif ctype == "code_inline":
            run = paragraph.add_run(child.content)
            run.font.name = _CODE_FONT
        elif ctype == "strong_open":
            bold = True
        elif ctype == "strong_close":
            bold = False
        elif ctype == "em_open":
            italic = True
        elif ctype == "em_close":
            italic = False
        elif ctype in ("s_open",):
            strike = True
        elif ctype in ("s_close",):
            strike = False
        elif ctype == "softbreak":
            paragraph.add_run(" ")
        elif ctype == "hardbreak":
            paragraph.add_run().add_break()
        elif ctype == "image":
            _emit_image(document, paragraph, child, base_dir=base_dir)
        elif ctype in ("link_open", "link_close"):
            # Render link text inline; the URL itself is dropped to keep the
            # document clean (text content arrives via the text children).
            continue
        elif child.content:
            paragraph.add_run(child.content)


def _emit_image(document, paragraph, child, *, base_dir: Path) -> None:
    """Embed an inline image from disk; leave a marker if it cannot be embedded."""
    from docx.shared import Inches

    attrs = dict(getattr(child, "attrs", None) or {})
    src = str(attrs.get("src") or "")
    alt = (child.content or attrs.get("alt") or "").strip()
    label = src or alt or "image"

    path = _resolve_image_path(src, base_dir)
    if path is None:
        paragraph.add_run(f"[image missing: {label}]")
        return
    try:
        run = paragraph.add_run()
        picture = run.add_picture(str(path))
    except Exception:
        paragraph.add_run(f"[image missing: {path.name}]")
        return

    max_width = Inches(6)
    if picture.width and picture.width > max_width:
        ratio = max_width / picture.width
        picture.width = int(picture.width * ratio)
        picture.height = int(picture.height * ratio)


def _resolve_image_path(src: str, base_dir: Path) -> Path | None:
    if not src or "://" in src or src.startswith("data:"):
        return None  # remote / data URIs are not embedded from disk
    candidate = Path(src)
    if not candidate.is_absolute():
        candidate = (base_dir / candidate)
    try:
        candidate = candidate.resolve()
    except OSError:
        return None
    if not candidate.is_file():
        return None
    if candidate.suffix.lower() not in _IMAGE_EXTS:
        return None
    return candidate


def _emit_code_block(document, content: str) -> None:
    from docx.shared import Pt

    text = content.rstrip("\n")
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(18)
    _shade_paragraph(paragraph, _CODE_SHADE)
    run = paragraph.add_run(text)
    run.font.name = _CODE_FONT
    run.font.size = Pt(9.5)


def _shade_paragraph(paragraph, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def _shade_cell(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _emit_table(document, tokens, start: int, *, base_dir: Path) -> int:
    """Build a Word table from a markdown-it table token run. Returns next index."""
    rows: list[list] = []
    current: list | None = None
    i = start + 1
    n = len(tokens)
    while i < n and tokens[i].type != "table_close":
        ttype = tokens[i].type
        if ttype == "tr_open":
            current = []
        elif ttype == "tr_close":
            if current is not None:
                rows.append(current)
            current = None
        elif ttype in ("th_open", "td_open") and current is not None:
            inline = tokens[i + 1] if i + 1 < n and tokens[i + 1].type == "inline" else None
            current.append(inline)
        i += 1

    if rows:
        ncols = max(len(row) for row in rows)
        table = document.add_table(rows=0, cols=ncols)
        try:
            table.style = "Table Grid"  # gives every cell a visible border
        except KeyError:
            pass
        for row_index, row in enumerate(rows):
            cells = table.add_row().cells
            for col in range(ncols):
                inline = row[col] if col < len(row) else None
                paragraph = cells[col].paragraphs[0]
                _emit_inline(document, paragraph, inline, base_dir=base_dir)
                if row_index == 0:
                    _shade_cell(cells[col], _HEADER_SHADE)
                    for run in paragraph.runs:
                        run.bold = True

    return i + 1  # advance past table_close
