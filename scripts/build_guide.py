#!/usr/bin/env python3
from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
from pathlib import Path
from typing import List

import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
from openai import OpenAI
from rich.console import Console

from sanitize_markdown_math import sanitize_markdown_math

console = Console()


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return path.read_text(encoding="utf-8", errors="replace")

    if suffix == ".pdf":
        chunks = []
        with fitz.open(path) as doc:
            for i, page in enumerate(doc, start=1):
                text = page.get_text("text")
                if text.strip():
                    chunks.append(f"\n\n--- Page {i} ---\n{text}")
        return "\n".join(chunks)

    if suffix == ".docx":
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    if suffix == ".pptx":
        prs = Presentation(path)
        chunks = []
        for i, slide in enumerate(prs.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
            if parts:
                chunks.append(f"\n\n--- Slide {i} ---\n" + "\n".join(parts))
        return "\n".join(chunks)

    raise ValueError(f"Unsupported file type: {suffix}")


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def chunk_text(text: str, max_chars: int = 18000) -> List[str]:
    paragraphs = re.split(r"\n\s*\n", text)
    chunks: list[str] = []
    current: list[str] = []
    current_len = 0

    for p in paragraphs:
        p = p.strip()
        if not p:
            continue

        if len(p) > max_chars:
            if current:
                chunks.append("\n\n".join(current))
                current, current_len = [], 0
            for i in range(0, len(p), max_chars):
                chunks.append(p[i:i + max_chars])
            continue

        if current_len + len(p) > max_chars and current:
            chunks.append("\n\n".join(current))
            current = [p]
            current_len = len(p)
        else:
            current.append(p)
            current_len += len(p)

    if current:
        chunks.append("\n\n".join(current))

    return chunks


def read_prompt(name: str) -> str:
    return (Path(__file__).resolve().parents[1] / "prompts" / name).read_text(encoding="utf-8")


def client() -> OpenAI:
    base_url = os.getenv("LLM_BASE_URL", "http://127.0.0.1:8080/v1")
    api_key = os.getenv("LLM_API_KEY", "local")
    return OpenAI(base_url=base_url, api_key=api_key)


def call_llm(messages: list[dict], temperature: float = 0.2) -> str:
    model = os.getenv("LLM_MODEL", "local-model")
    c = client()
    resp = c.chat.completions.create(
        model=model,
        messages=messages,
        temperature=temperature,
    )
    return resp.choices[0].message.content or ""


def generate_chunk_guide(source: str, title: str, mode: str, chunk_id: int, total: int) -> str:
    system = read_prompt("study_guide_system.md")
    user_template = read_prompt("study_guide_user.md")
    chunk_note = f"\n\nThis is chunk {chunk_id}/{total}. Cover only this chunk carefully.\n"
    user = user_template.format(title=title, mode=mode, source=source + chunk_note)
    return call_llm([
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ])


def merge_guides(parts: list[str], title: str, mode: str) -> str:
    if len(parts) == 1:
        return parts[0]

    system = read_prompt("study_guide_system.md")
    joined = "\n\n--- PART BREAK ---\n\n".join(parts)
    user = f"""Merge these partial study guides into ONE final clean study guide.

Title: {title}
Mode: {mode}

Rules:
- Remove duplicated headings.
- Keep all important concepts.
- Keep formulas in correct Markdown math.
- Keep the same Claude-style simple explanation.
- Make the final output coherent and not repetitive.
- Keep solved examples and exam questions if present.

Partial guides:
{joined}
"""
    return call_llm([
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ])


def validate_math(path: Path, show_math: bool = False) -> None:
    validator = Path(__file__).with_name("validate_math.js")
    if not validator.exists():
        console.print("[yellow]validate_math.js not found; skipping validation[/yellow]")
        return
    cmd = ["node", str(validator), str(path)]
    if show_math:
        cmd.append("--show-math")
    proc = subprocess.run(cmd)
    if proc.returncode != 0:
        raise SystemExit(proc.returncode)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path, help="PDF, DOCX, PPTX, TXT, or MD")
    parser.add_argument("--title", required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--mode", default="exam", choices=["exam", "theory", "quick", "deep"])
    parser.add_argument("--chunk-chars", type=int, default=18000)
    parser.add_argument("--no-sanitize", action="store_true", help="Do not run Markdown math sanitizer.")
    parser.add_argument("--sanitize-only", action="store_true", help="Only extract/sanitize input text; do not call LLM.")
    parser.add_argument("--validate-math", action="store_true", help="Validate output math with KaTeX using Node.")
    parser.add_argument("--show-math", action="store_true", help="Show math expressions during validation.")
    args = parser.parse_args()

    args.out.parent.mkdir(parents=True, exist_ok=True)

    console.print(f"[bold]Extracting:[/bold] {args.input}")
    raw = clean_text(extract_text(args.input))

    if args.sanitize_only:
        out = sanitize_markdown_math(raw) if not args.no_sanitize else raw
        args.out.write_text(out, encoding="utf-8")
        console.print(f"[green]Saved sanitized Markdown:[/green] {args.out}")
        if args.validate_math:
            validate_math(args.out, args.show_math)
        return

    chunks = chunk_text(raw, args.chunk_chars)
    console.print(f"[bold]Chunks:[/bold] {len(chunks)}")

    parts = []
    for i, chunk in enumerate(chunks, start=1):
        console.print(f"[cyan]Generating chunk {i}/{len(chunks)}...[/cyan]")
        parts.append(generate_chunk_guide(chunk, args.title, args.mode, i, len(chunks)))

    console.print("[cyan]Merging/finalizing...[/cyan]")
    final = merge_guides(parts, args.title, args.mode)

    if not args.no_sanitize:
        final = sanitize_markdown_math(final)

    args.out.write_text(final, encoding="utf-8")
    console.print(f"[green]Saved:[/green] {args.out}")

    if args.validate_math:
        validate_math(args.out, args.show_math)


if __name__ == "__main__":
    main()
