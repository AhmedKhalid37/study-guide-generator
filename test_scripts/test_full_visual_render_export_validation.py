#!/usr/bin/env python3
"""Slice 91: full figure insertion export/render validation.

Run with:

    python test_scripts/test_full_visual_render_export_validation.py

No external APIs and no PyMuPDF/Tesseract/Mistral/Gemini/Chandra dependency.

Slice 90 added the full non-table figure insertion v2 path (behind
``GUIDEFORGE_ENABLE_FULL_VISUAL_INSERTION``), which can place *many* safe
``assets/<slug>.png`` figure refs into a guide's ``clean.md`` instead of the legacy
top one or two. Slice 91 validates — and where load-bearing, fixes — the
render/export asset path so a guide with many inserted figures rides through
Markdown / HTML / PDF / DOCX / export-ZIP safely.

The one load-bearing cap was the export bundle ride-along, which discarded every
referenced figure past the legacy hard bound of two. Slice 91 adds
``find_all_exportable_visual_assets`` (uncapped, only bounded by the same
defensive ``_FULL_INSERTION_HARD_CEILING`` that guards a pathological manifest) and
switches the export bundle to it, so ALL referenced safe figures ride along. The
HTML/PDF/DOCX renderers were already uncapped (they render every referenced
``assets/<slug>.png`` from the job dir); this suite proves that and proves missing /
unsafe refs degrade safely with no path leakage.

Synthetic only: tiny generated temp dirs and a few bytes of fake PNG. No real
image / PDF / DOCX / ZIP fixture is added to the repo.

Parts:
  * Part A — PURE discovery checks (always run): ``find_all_exportable_visual_assets``
    discovers all safe refs (not just two), the legacy capped helper still stops at
    two, unsafe/missing refs are rejected/degraded, and the defensive ceiling holds.
  * Part B — HTML render (runs when the Markdown renderer imports): all referenced
    figures survive into the HTML the PDF renderer consumes.
  * Part C — DOCX render (SKIPPED when python-docx is absent): every present figure
    embeds; a missing figure leaves a safe marker with no raw path.
  * Part D — export BUNDLE (SKIPPED when FastAPI is absent): all referenced safe
    figures ride along in the ZIP (not just two); unsafe/missing degrade; the bundle
    index records safe relative refs only and leaks nothing.
"""
from __future__ import annotations

import io
import json
import os
import re
import struct
import sys
import tempfile
import zlib
import zipfile
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from pipeline.visual_markdown_insertion import (  # noqa: E402
    _FULL_INSERTION_HARD_CEILING,
    _HARD_MAX_IMAGES,
    find_all_exportable_visual_assets,
    find_exportable_visual_pilot_assets,
)

PASS = 0
FAIL = 0

# Leak detectors. A safe generated ref like ``assets/figure-001.png`` is allowed;
# anything resembling a real path / url / credential / data-uri / base64 image body
# is not.
KEYLIKE = re.compile(r"(sk-|sk_)[A-Za-z0-9_\-]{16,}")
PATHLIKE = re.compile(r"(/home/|/usr/|/etc/|/var/|/tmp/|C:\\\\|\\\\\\\\)")
URLLIKE = re.compile(r"https?://")
DATAURI = re.compile(r"data:[a-z]+/[a-z0-9.+-]+;base64,", re.IGNORECASE)

# A recognisable ASCII marker embedded in the fake PNG bytes. It must appear in a
# zipped image entry (proving the ride-along works) but NEVER in user-facing
# metadata (the bundle index, HTML, warnings).
PNG_BODY_SENTINEL = "png_image_body_sentinel"
PNG_BYTES = b"\x89PNG\r\n\x1a\n" + PNG_BODY_SENTINEL.encode("ascii") + b" fake png"


def _real_png_bytes() -> bytes:
    """A minimal *valid* 1x1 PNG, generated at runtime (never committed to the repo).

    python-docx parses the PNG header to embed a picture, so the byte-marker fake
    above is enough for byte-copy ride-along checks but not for real embedding. This
    builds a tiny standards-valid PNG so Part C can prove the DOCX renderer embeds
    every present figure rather than silently dropping it.
    """
    def _chunk(typ: bytes, data: bytes) -> bytes:
        body = typ + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    signature = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)  # 1x1, 8-bit, truecolour
    idat = zlib.compress(b"\x00\xff\xff\xff")  # one scanline: filter byte + white pixel
    return signature + _chunk(b"IHDR", ihdr) + _chunk(b"IDAT", idat) + _chunk(b"IEND", b"")

# Safe generated refs in deterministic order — the only refs allowed to surface.
SAFE_SLUGS = [f"figure_{i:03d}" for i in range(1, 6)]  # 5 figures => proves "> 2"
SAFE_REFS = [f"assets/{slug}.png" for slug in SAFE_SLUGS]


def check(name: str, condition: bool, detail: str = "") -> None:
    global PASS, FAIL
    if condition:
        PASS += 1
        print(f"[PASS] {name}")
    else:
        FAIL += 1
        msg = f"[FAIL] {name}"
        if detail:
            msg += f" - {detail}"
        print(msg)


def _scan_for_leak(text: str) -> str | None:
    if KEYLIKE.search(text):
        return "credential-like value"
    if PATHLIKE.search(text):
        return "filesystem-path-like value"
    if URLLIKE.search(text):
        return "url-like value"
    if DATAURI.search(text):
        return "data-uri/base64-like value"
    if PNG_BODY_SENTINEL in text:
        return "image-body-bytes leak"
    return None


def _markdown_with_refs(refs: list[str]) -> str:
    """Synthetic guide Markdown referencing each safe ref as a standard image."""
    body = ["# Synthetic Guide", "", "Intro paragraph.", ""]
    for index, ref in enumerate(refs, start=1):
        body.append(f"![Source visual, page {index}.]({ref})")
        body.append("")
    body.append("Closing paragraph.")
    return "\n".join(body) + "\n"


def _fake_job(job_dir: Path) -> SimpleNamespace:
    """Minimal duck-typed job for the pure helpers (only ``dir`` + ``clean_md``)."""
    return SimpleNamespace(dir=job_dir, clean_md=job_dir / "clean.md")


def _write_guide(
    job_dir: Path, refs: list[str], present_slugs: list[str], *, png_bytes: bytes = PNG_BYTES
) -> None:
    """Write a synthetic clean.md plus the assets that should exist on disk."""
    (job_dir / "clean.md").write_text(_markdown_with_refs(refs), encoding="utf-8")
    assets = job_dir / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    for slug in present_slugs:
        (assets / f"{slug}.png").write_bytes(png_bytes)


# ---------------------------------------------------------------------------
# Part A — pure discovery checks (no FastAPI / no renderers)
# ---------------------------------------------------------------------------


def part_a() -> None:
    # --- all 5 safe referenced+present figures are discovered (not capped at 2)
    with tempfile.TemporaryDirectory() as tmp:
        job_dir = Path(tmp)
        _write_guide(job_dir, SAFE_REFS, SAFE_SLUGS)
        job = _fake_job(job_dir)

        found = find_all_exportable_visual_assets(job)
        check("discovers all five referenced figures (not top-2)", found == SAFE_REFS, str(found))
        check("discovery preserves deterministic reference order", found == SAFE_REFS, str(found))
        check("legacy helper still stops at the hard cap of two",
              find_exportable_visual_pilot_assets(job) == SAFE_REFS[:_HARD_MAX_IMAGES],
              str(find_exportable_visual_pilot_assets(job)))
        leak = _scan_for_leak(" ".join(found))
        check("discovered refs carry no path/url/credential/base64 leak", leak is None, leak or "")
        check("discovered refs are all safe generated assets/<slug>.png",
              all(re.fullmatch(r"assets/[A-Za-z0-9_]+\.png", r) for r in found), str(found))

    # --- duplicate references are de-duplicated, order preserved -------------
    with tempfile.TemporaryDirectory() as tmp:
        job_dir = Path(tmp)
        dup_refs = [SAFE_REFS[0], SAFE_REFS[1], SAFE_REFS[0], SAFE_REFS[2]]
        _write_guide(job_dir, dup_refs, SAFE_SLUGS[:3])
        found = find_all_exportable_visual_assets(_fake_job(job_dir))
        check("duplicate refs de-duplicated, first-seen order kept",
              found == [SAFE_REFS[0], SAFE_REFS[1], SAFE_REFS[2]], str(found))

    # --- a missing referenced file is skipped; present neighbours survive ----
    with tempfile.TemporaryDirectory() as tmp:
        job_dir = Path(tmp)
        # All five referenced, but slug 003 is NOT written to disk.
        present = [s for s in SAFE_SLUGS if s != "figure_003"]
        _write_guide(job_dir, SAFE_REFS, present)
        found = find_all_exportable_visual_assets(_fake_job(job_dir))
        check("missing referenced figure is skipped, others ride along",
              found == [r for r in SAFE_REFS if "figure_003" not in r], str(found))
        check("missing figure degrades with no path leakage",
              _scan_for_leak(" ".join(found)) is None)

    # --- every referenced file missing => empty, no crash, no leak ----------
    with tempfile.TemporaryDirectory() as tmp:
        job_dir = Path(tmp)
        _write_guide(job_dir, SAFE_REFS, present_slugs=[])
        found = find_all_exportable_visual_assets(_fake_job(job_dir))
        check("all-missing degrades to empty list", found == [], str(found))

    # --- unsafe refs (abs / traversal / url / data-uri / non-png / nested) ---
    with tempfile.TemporaryDirectory() as tmp:
        job_dir = Path(tmp)
        unsafe_md = "\n\n".join(
            [
                "# Guide",
                "![abs](/etc/passwd.png)",
                "![dotdot](assets/../secret.png)",
                "![url](https://example.com/a.png)",
                "![datauri](data:image/png;base64,AAAA.png)",
                "![nonpng](assets/real.jpg)",
                "![nested](assets/sub/real.png)",
                "![backslash](assets\\evil.png)",
            ]
        )
        (job_dir / "clean.md").write_text(unsafe_md + "\n", encoding="utf-8")
        assets = job_dir / "assets"
        assets.mkdir(parents=True, exist_ok=True)
        # Files that an unsafe ref *might* point at — must still never be selected.
        (assets / "secret.png").write_bytes(PNG_BYTES)
        (assets / "real.jpg").write_bytes(PNG_BYTES)
        found = find_all_exportable_visual_assets(_fake_job(job_dir))
        check("every unsafe ref is rejected (none discovered)", found == [], str(found))

    # --- defensive ceiling: a pathological manifest is bounded --------------
    with tempfile.TemporaryDirectory() as tmp:
        job_dir = Path(tmp)
        many_slugs = [f"figure_{i:04d}" for i in range(_FULL_INSERTION_HARD_CEILING + 25)]
        many_refs = [f"assets/{s}.png" for s in many_slugs]
        _write_guide(job_dir, many_refs, many_slugs)
        found = find_all_exportable_visual_assets(_fake_job(job_dir))
        check("pathological manifest is bounded by the defensive ceiling",
              len(found) == _FULL_INSERTION_HARD_CEILING, str(len(found)))

    # --- explicit limit override is honoured (and clamped) ------------------
    with tempfile.TemporaryDirectory() as tmp:
        job_dir = Path(tmp)
        _write_guide(job_dir, SAFE_REFS, SAFE_SLUGS)
        job = _fake_job(job_dir)
        check("limit=3 yields exactly three refs",
              find_all_exportable_visual_assets(job, limit=3) == SAFE_REFS[:3])
        check("limit=0 yields no refs",
              find_all_exportable_visual_assets(job, limit=0) == [])
        check("bad limit falls back to the ceiling (still all five)",
              find_all_exportable_visual_assets(job, limit="oops") == SAFE_REFS)

    # --- non-existent clean.md => empty, never raises -----------------------
    with tempfile.TemporaryDirectory() as tmp:
        job_dir = Path(tmp)
        check("no clean.md degrades to empty list",
              find_all_exportable_visual_assets(_fake_job(job_dir)) == [])


# ---------------------------------------------------------------------------
# Part B — HTML render preserves every referenced figure
# ---------------------------------------------------------------------------


def part_b() -> None:
    try:
        from pipeline.html_renderer import render_markdown
    except Exception as exc:  # markdown renderer deps not importable
        print(f"[SKIP] HTML render section ({type(exc).__name__})")
        return

    markdown = _markdown_with_refs(SAFE_REFS)
    html = render_markdown(markdown, title="Synthetic Guide")
    for ref in SAFE_REFS:
        check(f"HTML preserves {ref}", ref in html)
    check("HTML emits one <img> per referenced figure (none dropped)",
          html.count("<img") == len(SAFE_REFS), str(html.count("<img")))
    check("HTML carries no path/url/credential/base64 leak",
          _scan_for_leak(html) is None)


# ---------------------------------------------------------------------------
# Part C — DOCX render embeds every present figure / safe-marks the missing one
# ---------------------------------------------------------------------------


def part_c() -> None:
    try:
        import docx  # noqa: F401

        from pipeline.docx_renderer import render_docx
    except Exception as exc:  # python-docx not installed in host python
        print(f"[SKIP] DOCX render section ({type(exc).__name__})")
        return

    real_png = _real_png_bytes()

    # --- all five present figures embed without loss ------------------------
    with tempfile.TemporaryDirectory() as tmp:
        job_dir = Path(tmp)
        _write_guide(job_dir, SAFE_REFS, SAFE_SLUGS, png_bytes=real_png)
        out = job_dir / "final.docx"
        render_docx(job_dir / "clean.md", out, title="Synthetic Guide")
        check("DOCX render of five figures produces a file", out.is_file())
        document = docx.Document(str(out))
        # python-docx exposes embedded images as inline shapes.
        shape_count = len(document.inline_shapes)
        check("DOCX embeds all five present figures (none dropped)",
              shape_count == len(SAFE_REFS), str(shape_count))
        body_text = "\n".join(p.text for p in document.paragraphs)
        check("DOCX body has no image-missing marker when all present",
              "[image missing" not in body_text, body_text[:80])
        check("DOCX body carries no path/url/credential/base64 leak",
              _scan_for_leak(body_text) is None)

    # --- a missing figure leaves a safe marker, no raw path -----------------
    with tempfile.TemporaryDirectory() as tmp:
        job_dir = Path(tmp)
        present = [s for s in SAFE_SLUGS if s != "figure_003"]
        _write_guide(job_dir, SAFE_REFS, present, png_bytes=real_png)
        out = job_dir / "final.docx"
        render_docx(job_dir / "clean.md", out, title="Synthetic Guide")
        document = docx.Document(str(out))
        body_text = "\n".join(p.text for p in document.paragraphs)
        check("DOCX marks the missing figure rather than crashing",
              "[image missing" in body_text, body_text[:120])
        check("DOCX missing-figure marker uses only the safe ref (no real path)",
              _scan_for_leak(body_text) is None, body_text[:120])
        check("DOCX still embeds the four present figures",
              len(document.inline_shapes) == len(present), str(len(document.inline_shapes)))


# ---------------------------------------------------------------------------
# Part D — export bundle rides ALL referenced figures along (needs FastAPI)
# ---------------------------------------------------------------------------


def part_d() -> int:
    try:
        from fastapi import HTTPException

        from api import server  # noqa: WPS433
    except Exception as exc:  # FastAPI / deps not importable in host python
        print(f"[SKIP] export bundle section ({type(exc).__name__})")
        return 0

    from pipeline.job_manager import Job

    def _make_job(
        root: Path,
        job_id: str,
        *,
        with_pdf: bool = True,
        clean_md: str | None = None,
        assets: dict[str, bytes] | None = None,
    ) -> Job:
        job = Job(id=job_id, root=root)
        job.input_dir.mkdir(parents=True, exist_ok=True)
        job.logs_dir.mkdir(parents=True, exist_ok=True)
        job._write_manifest({"id": job.id, "status": "completed", "title": job.id})
        if with_pdf:
            job.final_pdf.write_bytes(b"%PDF-1.4 fake pdf bytes")
        if assets:
            job.assets_dir.mkdir(parents=True, exist_ok=True)
            for name, data in assets.items():
                (job.assets_dir / name).write_bytes(data)
        if clean_md is not None:
            job.save_text(job.clean_md, clean_md)
        return job

    def _bundle(jobs_by_id, job_ids, artifacts):
        original = server._get_job

        def fake_get_job(raw_id):
            jid = str(raw_id)
            if jid in jobs_by_id:
                return jobs_by_id[jid]
            raise HTTPException(status_code=404, detail="Job not found.")

        server._get_job = fake_get_job
        try:
            resp = server.export_bundle(
                server.BundleRequest(job_ids=job_ids, artifacts=artifacts)
            )
        finally:
            server._get_job = original
        return resp

    def _zip(resp):
        return zipfile.ZipFile(io.BytesIO(resp.body))

    def _names(resp) -> list[str]:
        with _zip(resp) as zf:
            return zf.namelist()

    clean_five = _markdown_with_refs(SAFE_REFS)

    # --- all five referenced figures ride along (proves cap-2 is gone) ------
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        job = _make_job(
            root, "job-five",
            clean_md=clean_five,
            assets={f"{slug}.png": PNG_BYTES for slug in SAFE_SLUGS},
        )
        resp = _bundle({"job-five": job}, ["job-five"], ["pdf", "markdown"])
        names = _names(resp)
        png_entries = [n for n in names if n.lower().endswith(".png")]
        check("bundle rides ALL five referenced figures (not two)",
              len(png_entries) == len(SAFE_SLUGS), str(png_entries))
        check("each safe referenced figure is present in the bundle",
              all(any(n.endswith(f"/{ref}") for n in names) for ref in SAFE_REFS), str(names))
        check("bundled figure entries are safe relative paths under the base dir",
              all(not p.startswith("/") and ".." not in p and "\\" not in p
                  and "/assets/" in p for p in png_entries), str(png_entries))
        with _zip(resp) as zf:
            manifest_text = zf.read("manifest.json").decode("utf-8")
        entry = json.loads(manifest_text)["jobs"][0]
        check("manifest lists all five safe refs under visual_pilot_assets",
              entry.get("visual_pilot_assets") == SAFE_REFS, str(entry.get("visual_pilot_assets")))
        check("manifest keeps the backward-compatible first-ref field",
              entry.get("visual_pilot_asset") == SAFE_REFS[0], str(entry.get("visual_pilot_asset")))
        check("bundle index leaks no image bytes / path / url / base64",
              _scan_for_leak(manifest_text) is None)
        # files_included counts only requested artifacts (pdf+markdown=2), never figures.
        check("ride-along figures never inflate files_included",
              json.loads(manifest_text)["files_included"] == 2,
              str(json.loads(manifest_text)["files_included"]))

    # --- a missing referenced figure is skipped; the rest still ride along --
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        present = {f"{s}.png": PNG_BYTES for s in SAFE_SLUGS if s != "figure_003"}
        job = _make_job(root, "job-miss", clean_md=clean_five, assets=present)
        resp = _bundle({"job-miss": job}, ["job-miss"], ["pdf"])
        names = _names(resp)
        png_entries = [n for n in names if n.lower().endswith(".png")]
        check("missing figure skipped; the other four still ride along",
              len(png_entries) == len(present), str(png_entries))
        check("the missing figure is absent from the bundle",
              not any(n.endswith("/assets/figure_003.png") for n in names), str(names))
        with _zip(resp) as zf:
            entry = json.loads(zf.read("manifest.json").decode("utf-8"))["jobs"][0]
        check("manifest records only the present safe refs",
              entry.get("visual_pilot_assets") == [r for r in SAFE_REFS if "figure_003" not in r],
              str(entry.get("visual_pilot_assets")))

    # --- unsafe refs are never bundled even if a matching file exists -------
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        unsafe_clean = "\n\n".join(
            [
                "# Guide",
                "![abs](/etc/passwd.png)",
                "![dotdot](assets/../secret.png)",
                "![url](https://example.com/a.png)",
                "![datauri](data:image/png;base64,AAAA.png)",
                "![nonpng](assets/real.jpg)",
                "![nested](assets/sub/real.png)",
            ]
        )
        job = _make_job(
            root, "job-unsafe",
            clean_md=unsafe_clean,
            assets={"secret.png": PNG_BYTES, "real.jpg": PNG_BYTES},
        )
        resp = _bundle({"job-unsafe": job}, ["job-unsafe"], ["pdf"])
        names = _names(resp)
        check("no unsafe-referenced file is ever bundled",
              not any(n.lower().endswith((".png", ".jpg")) for n in names), str(names))
        with _zip(resp) as zf:
            entry = json.loads(zf.read("manifest.json").decode("utf-8"))["jobs"][0]
        check("manifest records no figure for an all-unsafe guide",
              not entry.get("visual_pilot_assets"), str(entry.get("visual_pilot_assets")))

    return 0


def main() -> int:
    part_a()
    part_b()
    part_c()
    part_d()
    print(f"\n{PASS} passed, {FAIL} failed")
    return 1 if FAIL else 0


if __name__ == "__main__":
    raise SystemExit(main())
